import csv
import os
from datetime import datetime
from io import BytesIO

from django.conf import settings as django_settings
from django.http import HttpResponse
from django.template.loader import render_to_string

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

STATIC_DIR   = os.path.join(django_settings.BASE_DIR, 'statics')
ISF_LOGO     = os.path.join(STATIC_DIR, 'indusserviseflow logo.png')
ORCHASP_LOGO = os.path.join(STATIC_DIR, 'orchasp logo.png')


def _file_url(path):
    from urllib.request import pathname2url
    return 'file:///' + pathname2url(path).lstrip('/')


def _render_simulation_html(simulation, result, trends):
    context = {
        'simulation':       simulation,
        'result':           result,
        'trends':           trends,
        'isf_logo_path':    _file_url(ISF_LOGO),
        'orchasp_logo_path':_file_url(ORCHASP_LOGO),
        'generated_at':     datetime.now().strftime('%d %b %Y, %I:%M %p'),
        'printed_on':       datetime.now().strftime('%d-%m-%Y %H:%M:%S'),
    }
    return render_to_string('report/simulation_report.html', context)


# ── PDF via WeasyPrint ────────────────────────────────────────────────────────

def generate_pdf_report(simulation, result, trends) -> HttpResponse:
    html = _render_simulation_html(simulation, result, trends)
    try:
        from weasyprint import HTML
        pdf_bytes = HTML(string=html).write_pdf()
    except ImportError:
        response = HttpResponse(html, content_type='text/html; charset=utf-8')
        response['X-PDF-Error'] = 'weasyprint not installed; returning HTML'
        return response

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="ISF_Simulation_{simulation.id}.pdf"'
    return response


# ── Excel via openpyxl ────────────────────────────────────────────────────────

def _header_row(ws, row_num: int, color: str):
    fill = PatternFill('solid', fgColor=color)
    for cell in ws[row_num]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center')


def generate_excel_report(simulation, result, trends) -> HttpResponse:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Simulation Report'

    ws.append(['ISF Queue Simulation Report'])
    ws['A1'].font = Font(bold=True, size=14, color='1B2A6B')
    ws.append([])

    ws.append(['Simulation Parameters'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
    ws.append(['Parameter', 'Value'])
    _header_row(ws, ws.max_row, '1B2A6B')
    for row in [
        ('Organization',        simulation.organization.organization_name),
        ('Arrival Probability', simulation.arrival_probability),
        ('Service Probability', simulation.service_probability),
        ('Number of Arrivals',  simulation.number_of_arrivals),
        ('Time Horizon (hrs)',  simulation.time_horizon),
        ('Status',              simulation.status.upper()),
        ('Created At',          str(simulation.created_at)[:19]),
    ]:
        ws.append(row)
    ws.append([])

    ws.append(['Simulation Results'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
    ws.append(['Metric', 'Value'])
    _header_row(ws, ws.max_row, '0d9488')
    for row in [
        ('Total Customers',             result.total_customers),
        ('Customers Served',            result.customers_served),
        ('Currently Waiting',           result.currently_waiting),
        ('Left Without Service',        result.left_without_service),
        ('Average Waiting Time (min)',  round(result.average_waiting_time, 2)),
        ('Maximum Waiting Time (min)',  round(result.maximum_waiting_time, 2)),
        ('Idle Time (hrs)',             round(result.idle_time, 2)),
        ('Utilization (%)',             round(result.utilization, 2)),
        ('Recommended Doctors',         result.recommended_doctors),
        ('Recommendation',              result.recommendation_message),
    ]:
        ws.append(row)
    ws.append([])

    if trends:
        ws.append(['Hourly Trend Analysis'])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
        ws.append(['Hour', 'Avg Wait (min)', 'Max Queue Length', 'Utilization (%)'])
        _header_row(ws, ws.max_row, 'f97316')
        for tr in trends:
            ws.append([tr.hour, round(tr.average_wait, 2), tr.maximum_queue, round(tr.utilization, 2)])

    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 18

    buffer = BytesIO()
    wb.save(buffer)
    file_bytes = buffer.getvalue()

    response = HttpResponse(
        file_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="ISF_Simulation_{simulation.id}.xlsx"'
    response['Content-Length'] = len(file_bytes)
    return response


# ── CSV ───────────────────────────────────────────────────────────────────────

def generate_csv_report(simulation, result, trends) -> HttpResponse:
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="ISF_Simulation_{simulation.id}.csv"'
    writer = csv.writer(response)

    writer.writerow(['ISF Queue Simulation Report'])
    writer.writerow([])
    writer.writerow(['Simulation Parameters'])
    writer.writerow(['Parameter', 'Value'])
    for row in [
        ('Organization',        simulation.organization.organization_name),
        ('Arrival Probability', simulation.arrival_probability),
        ('Service Probability', simulation.service_probability),
        ('Number of Arrivals',  simulation.number_of_arrivals),
        ('Time Horizon (hrs)',  simulation.time_horizon),
        ('Status',              simulation.status.upper()),
        ('Created At',          str(simulation.created_at)[:19]),
    ]:
        writer.writerow(row)

    writer.writerow([])
    writer.writerow(['Simulation Results'])
    writer.writerow(['Metric', 'Value'])
    for row in [
        ('Total Customers',             result.total_customers),
        ('Customers Served',            result.customers_served),
        ('Currently Waiting',           result.currently_waiting),
        ('Left Without Service',        result.left_without_service),
        ('Average Waiting Time (min)',  round(result.average_waiting_time, 2)),
        ('Maximum Waiting Time (min)',  round(result.maximum_waiting_time, 2)),
        ('Idle Time (hrs)',             round(result.idle_time, 2)),
        ('Utilization (%)',             round(result.utilization, 2)),
        ('Recommended Doctors',         result.recommended_doctors),
        ('Recommendation',              result.recommendation_message),
    ]:
        writer.writerow(row)

    if trends:
        writer.writerow([])
        writer.writerow(['Hourly Trend Analysis'])
        writer.writerow(['Hour', 'Avg Wait (min)', 'Max Queue Length', 'Utilization (%)'])
        for tr in trends:
            writer.writerow([tr.hour, round(tr.average_wait, 2), tr.maximum_queue, round(tr.utilization, 2)])

    return response
