import csv
import os
from datetime import datetime, time, timedelta

from django.conf import settings as django_settings
from django.db.models import Sum, Count, Q, Avg
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from appointments.models import Appointment, AppointmentService
from audit_logs.models import AuditLog
from employees.models import Employee
from organizations.models import Organization
from audit_logs.utils import log_action
from simulation.models import Simulation
from subscriptions.models import Subscription
from users.models import User


# ── Static asset paths ────────────────────────────────────────────────────────
STATIC_DIR   = os.path.join(django_settings.BASE_DIR, 'statics')
ISF_LOGO     = os.path.join(STATIC_DIR, 'indusserviseflow logo.png')
ORCHASP_LOGO = os.path.join(STATIC_DIR, 'orchasp logo.png')


def _file_url(path):
    from urllib.request import pathname2url
    return 'file:///' + pathname2url(path).lstrip('/')


SLOT_MAP = [
    ("8-10 AM",  8,  10),
    ("10-12 PM", 10, 12),
    ("12-2 PM",  12, 14),
    ("2-4 PM",   14, 16),
    ("4-6 PM",   16, 18),
    ("6-8 PM",   18, 20),
]


# ── Shared helpers ─────────────────────────────────────────────────────────────

def _get_appointment_summary(org_id, start_date=None, end_date=None):
    qs = Appointment.objects.filter(org_id=org_id)
    if start_date:
        qs = qs.filter(date__gte=start_date)
    if end_date:
        qs = qs.filter(date__lte=end_date)
    total = qs.count()
    status_counts = qs.values("status").annotate(count=Count("appointment_id"))
    status_map = {item["status"]: item["count"] for item in status_counts}
    return {
        "total_appointments": total,
        "status_breakdown": [
            {"status": s, "count": status_map.get(s, 0)}
            for s in ["Confirmed", "Waiting", "In Progress",
                      "Completed", "Cancelled", "No Show", "Left Queue"]
        ],
    }


def _get_peak_booking_hours(org_id, start_date=None, end_date=None):
    qs = Appointment.objects.filter(org_id=org_id)
    if start_date:
        qs = qs.filter(date__gte=start_date)
    if end_date:
        qs = qs.filter(date__lte=end_date)
    qs = (
        qs
        .values("time")
        .annotate(count=Count("appointment_id"))
        .order_by("-count")
    )
    return [{"hour": str(row["time"])[:5], "count": row["count"]} for row in qs]


def _get_staff_performance(org_id, start_date=None, end_date=None):
    qs = AppointmentService.objects.filter(appointment__org_id=org_id)
    if start_date:
        qs = qs.filter(appointment__date__gte=start_date)
    if end_date:
        qs = qs.filter(appointment__date__lte=end_date)
    qs = (
        qs
        .values("employee__employee_name")
        .annotate(
            total_assigned=Count("appointment_service_id"),
            completed=Count("appointment_service_id", filter=Q(status="Completed")),
            cancelled=Count("appointment_service_id", filter=Q(status="Cancelled")),
        )
        .order_by("employee__employee_name")
    )
    return [
        {
            "employee": row["employee__employee_name"] or "Unassigned",
            "total_assigned": row["total_assigned"],
            "completed": row["completed"],
            "cancelled": row["cancelled"],
            "completion_rate": round(
                (row["completed"] / row["total_assigned"] * 100), 1
            ) if row["total_assigned"] else 0.0,
        }
        for row in qs
    ]


def _get_wait_time_trend(org_id, start_date, end_date):
    """
    Average customer wait time (minutes) per day across the selected range,
    based on AppointmentService.waiting_time_min. Days with no recorded wait
    times are simply omitted rather than shown as a misleading zero.
    """
    qs = AppointmentService.objects.filter(
        appointment__org_id=org_id,
        appointment__date__range=(start_date, end_date),
        waiting_time_min__isnull=False,
    )
    rows = (
        qs.values("appointment__date")
        .annotate(avg_wait=Avg("waiting_time_min"))
        .order_by("appointment__date")
    )
    return [
        {
            "date": row["appointment__date"].strftime("%d %b"),
            "avg_wait_min": round(row["avg_wait"], 1) if row["avg_wait"] is not None else 0,
        }
        for row in rows
    ]


def _get_service_probability(org_id):
    sim = Simulation.objects.filter(
        organization_id=org_id
    ).order_by("-created_at").first()
    if not sim:
        return {
            "arrival_probability": None,
            "service_probability": None,
            "simulation_id": None,
            "simulation_status": None,
        }
    return {
        "arrival_probability": sim.arrival_probability,
        "service_probability": sim.service_probability,
        "simulation_id": sim.id,
        "simulation_status": sim.status,
    }


def _get_dashboard_analytics(org_id, start_date, end_date):
    qs = Appointment.objects.filter(org_id=org_id, date__range=(start_date, end_date))
    total = qs.count()
    served = qs.filter(status="Completed").count()

    # Active queue / queue waiting are "right now" counters, not a total to
    # sum over the report's selected date range. Summing over the range
    # would count any appointment still sitting in "Waiting"/"In Progress"
    # from earlier in the period as currently active, even if it's stale
    # and nobody is actually waiting today. Scope them to today's live
    # appointments instead, same as the org dashboard's live queue widget.
    live_qs = Appointment.objects.filter(org_id=org_id, date=timezone.localdate())
    active_queue = live_qs.filter(status__in=["Waiting", "In Progress"]).count()
    queue_waiting = live_qs.filter(status="Waiting").count()

    total_employees = Employee.objects.filter(org_id=org_id, status="Active").count()
    busy_employees = (
        qs.filter(status="In Progress")
        .values("services__employee")
        .distinct()
        .count()
    )
    utilization = round((busy_employees / total_employees * 100), 1) if total_employees else 0.0
    peak_hours = _get_peak_booking_hours(org_id, start_date, end_date)
    peak_hour = peak_hours[0]["hour"] if peak_hours else None
    return {
        "patients_today": total,
        "patients_served": served,
        "active_queue": active_queue,
        "queue_waiting": queue_waiting,
        "employee_utilization": utilization,
        "peak_hour": peak_hour,
    }


def _get_customer_flow(org_id, start_date, end_date):
    result = []
    cursor = start_date
    span = (end_date - start_date).days
    if span <= 7:
        for label, start, end in SLOT_MAP:
            booked = Appointment.objects.filter(org_id=org_id, date__range=(start_date, end_date), time__gte=f"{start:02d}:00", time__lt=f"{end:02d}:00").count()
            served = Appointment.objects.filter(org_id=org_id, date__range=(start_date, end_date), time__gte=f"{start:02d}:00", time__lt=f"{end:02d}:00", status="Completed").count()
            result.append({"slot": label, "booked": booked, "served": served, "pending": max(0, booked - served)})
        return result
    while cursor <= end_date:
        booked = Appointment.objects.filter(org_id=org_id, date=cursor).count()
        served = Appointment.objects.filter(org_id=org_id, date=cursor, status="Completed").count()
        result.append({"slot": cursor.strftime("%d %b"), "booked": booked, "served": served, "pending": max(0, booked - served)})
        cursor += timedelta(days=1)
    return result


def _parse_date_param(value):
    """Parse a 'YYYY-MM-DD' query param into a date, or None if missing/invalid."""
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def _resolve_report_date_range(start_date, end_date):
    """
    Fill in sensible defaults (last 30 days) when either bound is missing,
    and swap them if they were supplied in the wrong order. Returns
    (start_date, end_date, range_start_dt, range_end_dt) where the *_dt
    values are tz-aware datetimes covering the whole of each day.
    """
    today = timezone.localdate()

    if end_date is None:
        end_date = today
    if start_date is None:
        start_date = end_date - timedelta(days=30)
    if start_date > end_date:
        start_date, end_date = end_date, start_date

    range_start = timezone.make_aware(datetime.combine(start_date, time.min))
    range_end = timezone.make_aware(datetime.combine(end_date, time.max))

    return start_date, end_date, range_start, range_end


def _month_span(start_date, end_date, max_months=24, min_months=1):
    """
    Ordered list of (year, month) tuples covering start_date..end_date inclusive.

    Trend/growth charts need several buckets to draw a meaningful line, but a
    report's selected date range (e.g. the default "last 30 days") can collapse
    to a single calendar month. When `min_months` is greater than the number of
    months the caller's range actually spans, this walks the start back far
    enough (ending at end_date) to guarantee at least that many buckets, so
    growth charts don't silently render as one flat bar.
    """
    end_cursor = end_date.replace(day=1)
    cursor = start_date.replace(day=1)

    if min_months > 1:
        earliest_cursor = end_cursor
        for _ in range(min_months - 1):
            if earliest_cursor.month == 1:
                earliest_cursor = earliest_cursor.replace(year=earliest_cursor.year - 1, month=12)
            else:
                earliest_cursor = earliest_cursor.replace(month=earliest_cursor.month - 1)
        if earliest_cursor < cursor:
            cursor = earliest_cursor

    months = []
    while cursor <= end_cursor and len(months) < max_months:
        months.append((cursor.year, cursor.month))
        if cursor.month == 12:
            cursor = cursor.replace(year=cursor.year + 1, month=1)
        else:
            cursor = cursor.replace(month=cursor.month + 1)
    return months


def _week_span(start_date, end_date, max_weeks=12, min_weeks=1):
    """
    Ordered list of (week_start, week_end) date tuples covering start_date..end_date.

    `min_weeks` guarantees a minimum number of buckets (see `_month_span` above)
    by walking the start back from end_date when the selected range is too narrow.
    """
    if min_weeks > 1:
        earliest_start = end_date - timedelta(days=7 * min_weeks - 1)
        if earliest_start < start_date:
            start_date = earliest_start

    weeks = []
    cursor = start_date
    while cursor <= end_date and len(weeks) < max_weeks:
        week_end = min(cursor + timedelta(days=6), end_date)
        weeks.append((cursor, week_end))
        cursor = week_end + timedelta(days=1)
    return weeks


def _month_bounds(year, month):
    """
    Returns (start, end) as tz-aware datetimes covering the whole calendar
    month [start, end) in the app's local time zone.

    We deliberately compare against this range with >=/< instead of using
    created_on__year=/__month= lookups. With USE_TZ=True and a non-UTC
    TIME_ZONE (this project runs MySQL with TIME_ZONE="Asia/Kolkata"), those
    lookups need MySQL's CONVERT_TZ() to translate the stored UTC value into
    the local time zone before extracting the year/month. If MySQL's time
    zone tables haven't been loaded (a common setup gap — see
    `mysql_tzinfo_to_sql`), CONVERT_TZ() returns NULL and the __year=/__month=
    filter silently matches zero rows instead of raising an error. A plain
    datetime range comparison needs no such conversion and works reliably
    regardless of whether the MySQL time zone tables are loaded.
    """
    start = timezone.make_aware(datetime(year, month, 1))
    if month == 12:
        end = timezone.make_aware(datetime(year + 1, 1, 1))
    else:
        end = timezone.make_aware(datetime(year, month + 1, 1))
    return start, end


def _get_superadmin_report_data(start_date=None, end_date=None):
    start_date, end_date, range_start, range_end = _resolve_report_date_range(
        start_date, end_date
    )

    orgs_qs = Organization.objects.filter(is_deleted=False)
    orgs_in_range = orgs_qs.filter(created_on__range=(range_start, range_end))

    # The summary cards (Total/Active/Pending) are snapshot counts of the
    # platform's current state, matching the Organizations page and the
    # dashboard — they intentionally do NOT get scoped down to orgs_in_range.
    # Scoping them to the selected date range previously made this card only
    # count organizations *created* inside that window, which silently
    # excluded every older organization and made the number disagree with
    # the Organizations page (e.g. 12 here vs 22 there for the same data).
    # `orgs_in_range` is still used below for the period-specific breakdowns
    # (category mix, etc.) where "created within the selected range" is the
    # correct meaning.
    total_organizations = orgs_qs.count()
    active_organizations = orgs_qs.filter(status__iexact="active").count()
    pending_requests = orgs_qs.filter(status__iexact="pending").count()

    subs_in_range = Subscription.objects.filter(created_on__range=(range_start, range_end))
    monthly_revenue = (
        subs_in_range.filter(status__iexact="active")
        .aggregate(total=Sum("monthly_revenue"))["total"] or 0
    )
    # Same fix as above: "With Subscriptions" is a current-state snapshot
    # ("have an active plan"), not "subscribed within the selected range",
    # so count against all active subscriptions rather than subs_in_range.
    organizations_with_subscription = (
        Subscription.objects.filter(status__iexact="active")
        .values("organization").distinct().count()
    )

    # Growth/trend charts intentionally look at the FULL history (orgs_qs /
    # Subscription.objects.all()), not orgs_in_range / subs_in_range, and use a
    # guaranteed minimum span (min_months / min_weeks) rather than whatever the
    # report's date filter happens to be. Otherwise, with the default "last 30
    # days" window, these charts collapse to a single bucket that only counts
    # records created inside that same narrow window — i.e. an empty/flat
    # "growth" chart even when older organizations and subscriptions exist.
    subs_all = Subscription.objects.all()

    org_growth = []
    for year, month in _month_span(start_date, end_date, min_months=6):
        bucket_start, bucket_end = _month_bounds(year, month)
        org_growth.append({
            "month": datetime(year, month, 1).strftime("%b"),
            "organizations": orgs_qs.filter(
                created_on__gte=bucket_start, created_on__lt=bucket_end
            ).count(),
        })

    mrr_growth = []
    for idx, (week_start, week_end) in enumerate(
        _week_span(start_date, end_date, min_weeks=8), start=1
    ):
        week_end_dt = timezone.make_aware(datetime.combine(week_end, time.max))
        mrr_growth.append({
            "week": f"W{idx}",
            "mrr": (
                subs_all.filter(
                    status__iexact="active", created_on__lte=week_end_dt
                ).aggregate(total=Sum("monthly_revenue"))["total"] or 0
            ),
            "tenants": orgs_qs.filter(created_on__lte=week_end_dt).count(),
        })

    category_mix = [
        {"category": row["category__category_name"] or "Uncategorized", "count": row["count"]}
        for row in orgs_in_range.values("category__category_name")
        .annotate(count=Count("id")).order_by("-count")
    ]

    plan_distribution = [
        {"plan": row["plan__plan_name"] or "Unknown", "count": row["count"]}
        for row in subs_in_range.filter(status__iexact="active")
        .values("plan__plan_name").annotate(count=Count("id")).order_by("-count")
    ]

    subscription_changes = []
    for year, month in _month_span(start_date, end_date, min_months=6):
        bucket_start, bucket_end = _month_bounds(year, month)
        subscription_changes.append({
            "month": datetime(year, month, 1).strftime("%b"),
            "new_signups": subs_all.filter(
                created_on__gte=bucket_start, created_on__lt=bucket_end
            ).count(),
            "upgrades": 0,
            "downgrades": 0,
        })

    users_in_range = User.objects.filter(created_on__range=(range_start, range_end))
    audit_logs_in_range = AuditLog.objects.filter(audit_date__range=(range_start, range_end))

    return {
        "date_range": {
            "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(),
        },
        "cards": {
            "total_organizations": total_organizations,
            "active_organizations": active_organizations,
            "pending_requests": pending_requests,
            "monthly_revenue": monthly_revenue,
            "organizations_with_subscription": organizations_with_subscription,
        },
        "organization_growth": org_growth,
        "mrr_growth": mrr_growth,
        "category_mix": category_mix,
        "plan_distribution": plan_distribution,
        "subscription_changes": subscription_changes,
        "organization_status_breakdown": [
            {"status": item["status"], "count": item["count"]}
            for item in orgs_in_range.values("status").annotate(count=Count("id"))
        ],
        "subscription_status_breakdown": [
            {"status": item["status"], "count": item["count"]}
            for item in subs_in_range.values("status").annotate(count=Count("id"))
        ],
        "user_role_breakdown": [
            {"role": item["role"], "count": item["count"]}
            for item in users_in_range.values("role").annotate(count=Count("id"))
        ],
        "user_status_breakdown": [
            {"status": item["status"], "count": item["count"]}
            for item in users_in_range.values("status").annotate(count=Count("id"))
        ],
        "recent_users": list(
            users_in_range.order_by("-created_on")
            .values("username", "email", "role", "status", "created_on")[:20]
        ),
        "recent_audit_logs": list(
            audit_logs_in_range.order_by("-audit_date")
            .values("action_name", "action_screen", "username",
                    "role", "target_info", "audit_date")[:50]
        ),
    }


# ── HTML/PDF rendering via WeasyPrint ─────────────────────────────────────────

def _render_html(template_name, context):
    context.update({
        "isf_logo_path":     _file_url(ISF_LOGO),
        "orchasp_logo_path": _file_url(ORCHASP_LOGO),
        "printed_on":        datetime.now().strftime("%d-%m-%Y %H:%M:%S"),
        "generated_at":      datetime.now().strftime("%d %b %Y, %I:%M %p"),
    })
    return render_to_string(template_name, context)


def _html_to_pdf_response(html_string, filename):
    try:
        from weasyprint import HTML
        pdf_bytes = HTML(string=html_string).write_pdf()
    except ImportError:
        return Response(
            {
                "success": False,
                "message": "PDF export is unavailable because WeasyPrint is not installed.",
            },
            status=503,
        )
    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ── Flat-row builders (CSV / Excel) ───────────────────────────────────────────

def _build_report_rows(report_data):
    rows = []
    rows.append(["Organization Report"])
    rows.append([])
    rows.append(["Organization", report_data["organization"]["name"]])
    rows.append(["Category",    report_data["organization"]["category"]])
    rows.append(["City",        report_data["organization"]["city"]])
    rows.append(["State",       report_data["organization"]["state"]])
    rows.append(["Country",     report_data["organization"]["country"]])
    rows.append(["Status",      report_data["organization"]["status"]])
    rows.append([])
    rows.append(["Total Appointments", report_data["appointment_summary"]["total_appointments"]])
    rows.append([])
    rows.append(["Status Breakdown"])
    rows.append(["Status", "Count"])
    for item in report_data["appointment_summary"]["status_breakdown"]:
        rows.append([item["status"], item["count"]])
    rows.append([])
    rows.append(["Peak Booking Hours"])
    rows.append(["Hour", "Count"])
    for item in report_data["peak_booking_hours"]:
        rows.append([item["hour"], item["count"]])
    rows.append([])
    rows.append(["Staff Performance"])
    rows.append(["Employee", "Assigned", "Completed", "Cancelled", "Completion Rate (%)"])
    for item in report_data["staff_performance"]:
        rows.append([
            item["employee"], item["total_assigned"],
            item["completed"], item["cancelled"], item["completion_rate"],
        ])
    rows.append([])
    rows.append(["Service Probability"])
    rows.append(["Arrival Probability",  report_data["service_probability"]["arrival_probability"]])
    rows.append(["Service Probability",  report_data["service_probability"]["service_probability"]])
    rows.append(["Simulation ID",        report_data["service_probability"]["simulation_id"]])
    rows.append(["Simulation Status",    report_data["service_probability"]["simulation_status"]])
    rows.append([])
    rows.append(["Dashboard Analytics"])
    rows.append(["Metric", "Value"])
    for key, value in report_data["dashboard_analytics"].items():
        rows.append([key.replace("_", " ").title(), value])
    rows.append([])
    rows.append(["Customer Flow"])
    rows.append(["Slot", "Booked", "Served"])
    for item in report_data["customer_flow"]:
        rows.append([item["slot"], item["booked"], item["served"]])
    rows.append([])
    rows.append(["Wait Time Trend"])
    rows.append(["Date", "Avg Wait (min)"])
    for item in report_data["wait_time_trend"]:
        rows.append([item["date"], item["avg_wait_min"]])
    return rows


def _build_superadmin_report_rows(report_data):
    rows = []
    rows.append(["Super Admin Report"])
    rows.append([
        "Period",
        f"{report_data['date_range']['start_date']} to {report_data['date_range']['end_date']}",
    ])
    rows.append([])
    rows.append(["Total Organizations",              report_data["cards"]["total_organizations"]])
    rows.append(["Active Organizations",             report_data["cards"]["active_organizations"]])
    rows.append(["Pending Organization Requests",    report_data["cards"]["pending_requests"]])
    rows.append(["Monthly Revenue",                  report_data["cards"]["monthly_revenue"]])
    rows.append(["Organizations With Subscriptions", report_data["cards"]["organizations_with_subscription"]])
    rows.append([])
    rows.append(["Organization Growth"])
    rows.append(["Month", "Organizations"])
    for item in report_data["organization_growth"]:
        rows.append([item["month"], item["organizations"]])
    rows.append([])
    rows.append(["MRR Growth"])
    rows.append(["Week", "MRR", "Tenants"])
    for item in report_data["mrr_growth"]:
        rows.append([item["week"], item["mrr"], item["tenants"]])
    rows.append([])
    rows.append(["Category Mix"])
    rows.append(["Category", "Count"])
    for item in report_data["category_mix"]:
        rows.append([item["category"], item["count"]])
    rows.append([])
    rows.append(["Plan Distribution"])
    rows.append(["Plan", "Count"])
    for item in report_data["plan_distribution"]:
        rows.append([item["plan"], item["count"]])
    rows.append([])
    rows.append(["Subscription Changes"])
    rows.append(["Month", "New Signups", "Upgrades", "Downgrades"])
    for item in report_data["subscription_changes"]:
        rows.append([item["month"], item["new_signups"], item["upgrades"], item["downgrades"]])
    rows.append([])
    rows.append(["Organization Status Breakdown"])
    rows.append(["Status", "Count"])
    for item in report_data["organization_status_breakdown"]:
        rows.append([item["status"], item["count"]])
    rows.append([])
    rows.append(["Subscription Status Breakdown"])
    rows.append(["Status", "Count"])
    for item in report_data["subscription_status_breakdown"]:
        rows.append([item["status"], item["count"]])
    rows.append([])
    rows.append(["User Role Breakdown"])
    rows.append(["Role", "Count"])
    for item in report_data["user_role_breakdown"]:
        rows.append([item["role"], item["count"]])
    rows.append([])
    rows.append(["User Status Breakdown"])
    rows.append(["Status", "Count"])
    for item in report_data["user_status_breakdown"]:
        rows.append([item["status"], item["count"]])
    rows.append([])
    rows.append(["Recent Users"])
    rows.append(["Username", "Email", "Role", "Status", "Created On"])
    for item in report_data["recent_users"]:
        rows.append([
            item["username"], item["email"], item["role"],
            item["status"], item["created_on"].strftime("%Y-%m-%d %H:%M:%S"),
        ])
    rows.append([])
    rows.append(["Recent Audit Logs"])
    rows.append(["Action", "Screen", "User", "Role", "Target", "Created On"])
    for item in report_data["recent_audit_logs"]:
        rows.append([
            item["action_name"], item["action_screen"], item["username"],
            item["role"], item["target_info"],
            item["audit_date"].strftime("%Y-%m-%d %H:%M:%S"),
        ])
    return rows


# ══════════════════════════════════════════════════════════════════════════════
# SUPER ADMIN REPORT
# ══════════════════════════════════════════════════════════════════════════════

@api_view(["GET"])
@permission_classes([IsAuthenticated])
def superadmin_report(request):
    user_role = (getattr(request.user, "role", None) or "").upper()
    if user_role != "SUPER_ADMIN":
        return Response(
            {"success": False, "message": "Only Super Admin can access this report."},
            status=403,
        )

    start_date = _parse_date_param(request.GET.get("start_date"))
    end_date = _parse_date_param(request.GET.get("end_date"))
    report_data = _get_superadmin_report_data(start_date, end_date)
    export_format = request.GET.get("export_format", "json").lower()

    if export_format == "json":
        return Response({"success": True, "data": report_data})

    if export_format == "html":
        log_action(request.user, "Export", "Reports")
        html = _render_html(
            "report/super_admin_report.html",
            {"data": report_data, "generated_by": request.user.username},
        )
        return HttpResponse(html, content_type="text/html; charset=utf-8")

    if export_format == "pdf":
        log_action(request.user, "Export", "Reports")
        html = _render_html(
            "report/super_admin_report.html",
            {"data": report_data, "generated_by": request.user.username},
        )
        return _html_to_pdf_response(html, "ISF_SuperAdmin_Report.pdf")

    if export_format == "csv":
        log_action(request.user, "Export", "Reports")
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="superadmin_report.csv"'
        writer = csv.writer(response)
        for row in _build_superadmin_report_rows(report_data):
            writer.writerow(row)
        return response

    if export_format == "excel":
        log_action(request.user, "Export", "Reports")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Super Admin Report"
        section_headers = {
            "Super Admin Report", "Organization Growth", "MRR Growth",
            "Category Mix", "Plan Distribution", "Subscription Changes",
            "Organization Status Breakdown", "Subscription Status Breakdown",
            "User Role Breakdown", "User Status Breakdown",
            "Recent Users", "Recent Audit Logs",
        }
        row_num = 1
        for row in _build_superadmin_report_rows(report_data):
            for col_num, value in enumerate(row, 1):
                cell = sheet.cell(row=row_num, column=col_num)
                cell.value = value
                if row and row[0] in section_headers:
                    cell.font = Font(bold=True)
            row_num += 1
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="superadmin_report.xlsx"'
        workbook.save(response)
        return response

    return Response(
        {"success": False, "message": "Invalid export format. Supported: json, html, pdf, csv, excel."},
        status=400,
    )


# ══════════════════════════════════════════════════════════════════════════════
# ORGANIZATION REPORT
# ══════════════════════════════════════════════════════════════════════════════

@api_view(["GET"])
@permission_classes([IsAuthenticated])
def organization_report(request, org_id):
    organization = Organization.objects.filter(pk=org_id, is_deleted=False).first()
    if not organization:
        return Response(
            {"success": False, "message": "Organization not found."},
            status=404,
        )

    start_date = _parse_date_param(request.GET.get("from_date"))
    end_date = _parse_date_param(request.GET.get("to_date"))
    start_date, end_date, _, _ = _resolve_report_date_range(start_date, end_date)

    report_data = {
        "date_range": {
            "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(),
        },
        "organization": {
            "id":       organization.id,
            "name":     organization.organization_name,
            "category": organization.category.category_name if organization.category else None,
            "email":    organization.email,
            "city":     organization.city,
            "state":    organization.state,
            "country":  organization.country,
            "status":   organization.status,
        },
        "appointment_summary":  _get_appointment_summary(org_id, start_date, end_date),
        "peak_booking_hours":   _get_peak_booking_hours(org_id, start_date, end_date),
        "staff_performance":    _get_staff_performance(org_id, start_date, end_date),
        "service_probability":  _get_service_probability(org_id),
        "dashboard_analytics":  _get_dashboard_analytics(org_id, start_date, end_date),
        "customer_flow":        _get_customer_flow(org_id, start_date, end_date),
        "wait_time_trend":      _get_wait_time_trend(org_id, start_date, end_date),
    }

    export_format = request.GET.get("export_format", "json").lower()

    if export_format == "json":
        return Response({"success": True, "data": report_data})

    if export_format == "html":
        log_action(request.user, "Export", "Organization Reports")
        html = _render_html("report/organization_report.html", {"data": report_data})
        return HttpResponse(html, content_type="text/html; charset=utf-8")

    if export_format == "pdf":
        log_action(request.user, "Export", "Organization Reports")
        html = _render_html("report/organization_report.html", {"data": report_data})
        return _html_to_pdf_response(html, f"ISF_Org_{org_id}_Report.pdf")

    if export_format == "csv":
        log_action(request.user, "Export", "Organization Reports")
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = (
            f'attachment; filename="organization_{org_id}_report.csv"'
        )
        writer = csv.writer(response)
        for row in _build_report_rows(report_data):
            writer.writerow(row)
        return response

    if export_format == "excel":
        log_action(request.user, "Export", "Organization Reports")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Organization Report"
        section_headers = {
            "Organization Report", "Status Breakdown", "Peak Booking Hours",
            "Staff Performance", "Service Probability",
            "Dashboard Analytics", "Customer Flow", "Wait Time Trend",
        }
        row_num = 1
        for row in _build_report_rows(report_data):
            for col_num, value in enumerate(row, 1):
                cell = sheet.cell(row=row_num, column=col_num)
                cell.value = value
                if row and row[0] in section_headers:
                    cell.font = Font(bold=True)
            row_num += 1
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            f'attachment; filename="organization_{org_id}_report.xlsx"'
        )
        workbook.save(response)
        return response

    return Response(
        {"success": False, "message": "Invalid export format. Supported: json, html, pdf, csv, excel."},
        status=400,
    )