import logging
import random
import string
import csv
from django.db import transaction

from datetime import date, timedelta

from django.shortcuts import get_object_or_404
from django.conf import settings
from django.core.mail import send_mail
from django.core.paginator import Paginator
from django.db.models import Q, Prefetch
from django.http import HttpResponse

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from rest_framework import status

from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .models import Organization
from .serializers import (
    OrganizationRegistrationSerializer,
    OrganizationSerializer,
)

from users.models import User
from plans.models import Plan
from subscriptions.models import Subscription
from .dashboard_service import DashboardService


def _with_latest_subscription(queryset):
    """
    Attaches each organization's subscriptions (newest first, matching
    Subscription.Meta.ordering) as `prefetched_subscriptions`, so
    OrganizationSerializer can read plan/payment info in one extra query
    total instead of one query per organization.
    """
    return queryset.prefetch_related(
        Prefetch(
            "subscriptions",
            queryset=Subscription.objects.select_related("plan").order_by(
                "-created_on"
            ),
            to_attr="prefetched_subscriptions",
        )
    )


from notifications.services import (
    notify_org_admins,
    notify_super_admin,
)

from audit_logs.utils import log_action

import os
import html
from datetime import datetime
from urllib.request import pathname2url
from django.conf import settings as django_settings

_BADGE_CLASS_MAP = {
    "active": "badge-active",
    "approved": "badge-active",
    "completed": "badge-completed",
    "confirmed": "badge-active",
    "create": "badge-active",
    "assign": "badge-completed",
    "update": "badge-completed",
    "status change": "badge-pending",
    "export": "badge-pending",
    "pending": "badge-pending",
    "pending payment": "badge-pending",
    "pending activation": "badge-pending",
    "expiring soon": "badge-pending",
    "waiting": "badge-pending",
    "in progress": "badge-pending",
    "inactive": "badge-inactive",
    "login": "badge-inactive",
    "logout": "badge-inactive",
    "view": "badge-inactive",
    "cancelled": "badge-cancelled",
    "locked": "badge-rejected",
    "rejected": "badge-rejected",
    "reject": "badge-rejected",
    "delete": "badge-rejected",
    "expired": "badge-expired",
    "no show": "badge-rejected",
    "left queue": "badge-rejected",
}


def _status_badge_class(value):
    key = (str(value) if value is not None else "").strip().lower()
    return _BADGE_CLASS_MAP.get(key, "badge-inactive")


def _kpi_cards_html(cards):
    """cards: list of dicts {label, value, sub, css}. Renders in rows of 4."""
    if not cards:
        return ""
    rows_html = ""
    for i in range(0, len(cards), 4):
        chunk = cards[i : i + 4]
        cells = "".join(f"""<div class="metric-card {c["css"]}">
                  <div class="card-label">{html.escape(str(c["label"]))}</div>
                  <div class="card-value">{html.escape(str(c["value"]))}</div>
                  <div class="card-sub">{html.escape(str(c.get("sub", "")))}</div>
                </div>""" for c in chunk)
        rows_html += f'<div class="cards-grid cards-grid-4 mt-2">{cells}</div>'
    return rows_html


def _data_table_rows_html(headers, rows, status_col_index=None):
    header_cells = "".join(f"<th>{html.escape(str(h))}</th>" for h in headers)

    body_rows = ""
    for row in rows:
        cells = ""
        for idx, val in enumerate(row):
            text = "" if val is None else str(val)
            if status_col_index is not None and idx == status_col_index:
                css_class = _status_badge_class(text)
                cells += f'<td><span class="badge {css_class}">{html.escape(text)}</span></td>'
            else:
                cells += f"<td>{html.escape(text)}</td>"
        body_rows += f"<tr>{cells}</tr>"

    if not rows:
        body_rows = (
            f'<tr><td colspan="{len(headers)}" '
            'style="text-align:center;color:#94A3B8;padding:10mm 0;">'
            "No records found</td></tr>"
        )

    return header_cells, body_rows


def _build_list_report_pdf(
    report_title,
    report_subtitle,
    summary_cards,
    table_section_title,
    headers,
    rows,
    status_col_index,
    total_count,
    generated_by_username,
    filename,
    highlight=None,
):
    """
    Builds a PDF export styled like the IndusServiceFlow Organization Report
    (navy header banner, teal accent, numbered sections, KPI summary cards,
    striped status-badge table, Orchasp footer) using WeasyPrint -- no
    separate template file needed.

    summary_cards: list of dicts {"label", "value", "sub", "css"} where css
        is one of card-blue / card-teal / card-orange / card-red /
        card-grey / card-green / card-yellow.
    status_col_index: index (0-based) into `headers`/`rows` to render as a
        colored status badge, or None to render every column as plain text.
    highlight: optional dict {"label", "value"} shown as a teal callout box
        above the table (e.g. a total-revenue figure).
    """

    static_dir = os.path.join(django_settings.BASE_DIR, "statics")
    isf_logo = "file:///" + pathname2url(
        os.path.join(static_dir, "indusserviseflow logo.png")
    ).lstrip("/")
    orchasp_logo = "file:///" + pathname2url(
        os.path.join(static_dir, "orchasp logo.png")
    ).lstrip("/")

    generated_at = datetime.now().strftime("%d %b %Y, %I:%M %p")
    printed_on = datetime.now().strftime("%d-%m-%Y %H:%M:%S")

    header_cells, body_rows = _data_table_rows_html(headers, rows, status_col_index)
    cards_html = _kpi_cards_html(summary_cards)

    highlight_html = ""
    if highlight:
        highlight_html = f"""
        <div class="recommendation-box mb-4">
          <div class="rec-label">{html.escape(str(highlight["label"]))}</div>
          <div class="rec-text" style="font-size:14pt;">{html.escape(str(highlight["value"]))}</div>
        </div>
        """

    html_string = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="UTF-8"/>
    <style>
      * {{ box-sizing: border-box; margin: 0; padding: 0; }}
      :root {{
        --navy:#1B2A6B; --blue:#1565C0; --light-blue:#E8F0FE; --teal:#0d9488;
        --light-teal:#E6F7F5; --orange:#f97316; --light-orange:#FFF3E0;
        --red-light:#FCE4EC; --grey-bg:#F5F7FA; --grey-line:#CBD5E1;
        --white:#FFFFFF; --dark-text:#0F172A; --mid-text:#475569;
        --light-text:#94A3B8; --green:#16a34a; --green-light:#DCFCE7;
        --yellow-light:#FFFBEB;
      }}
      html {{ font-size: 10pt; }}
      body {{ font-family:'Helvetica Neue',Helvetica,Arial,sans-serif; color:var(--dark-text); }}
      .page-content {{ padding: 8mm 18mm 0 18mm; }}
      .report-header {{
        background: var(--navy); padding: 5mm 18mm; display: flex;
        align-items: center; justify-content: space-between;
        border-bottom: 2.5px solid var(--teal);
      }}
      .header-logo-block {{ display: flex; align-items: center; gap: 10px; }}
      .header-logo {{ height: 14mm; width: auto; object-fit: contain; }}
      .header-brand {{ color: #fff; }}
      .header-brand .brand-name {{ font-size: 13pt; font-weight: 700; }}
      .header-brand .brand-tagline {{ font-size: 7.5pt; color: rgba(255,255,255,0.72); margin-top: 1px; }}
      .header-center {{ text-align: center; color: #fff; }}
      .header-center .report-type {{ font-size: 13pt; font-weight: 700; letter-spacing: 1px; text-transform: uppercase; }}
      .header-center .report-subtitle {{ font-size: 7.5pt; color: rgba(255,255,255,0.72); margin-top: 2px; }}
      .header-meta {{ text-align: right; color: rgba(255,255,255,0.85); font-size: 7.5pt; line-height: 1.7; }}
      .header-meta strong {{ display: block; font-size: 8.5pt; color: #fff; }}
      .report-title-block {{ padding: 6mm 0 4mm 0; border-bottom: 1px solid var(--grey-line); margin-bottom: 6mm; }}
      .report-title-block h1 {{ font-size: 16pt; font-weight: 700; color: var(--navy); margin-bottom: 2px; }}
      .report-meta-row {{ display: flex; gap: 20px; flex-wrap: wrap; font-size: 8pt; color: var(--mid-text); margin-top: 4px; }}
      .report-meta-row strong {{ color: var(--dark-text); }}
      .section {{ margin-bottom: 7mm; }}
      .section-title {{
        font-size: 10.5pt; font-weight: 700; color: var(--navy);
        padding: 3mm 0 2mm 0; border-bottom: 2px solid var(--teal);
        margin-bottom: 4mm; display: flex; align-items: center; gap: 6px;
      }}
      .section-title .section-num {{
        background: var(--navy); color: #fff; font-size: 8pt; font-weight: 700;
        width: 18px; height: 18px; border-radius: 50%; display: inline-flex;
        align-items: center; justify-content: center; flex-shrink: 0;
      }}
      .cards-grid {{ display: grid; gap: 4mm; margin-bottom: 2mm; }}
      .cards-grid-4 {{ grid-template-columns: repeat(4, 1fr); }}
      .cards-grid-2 {{ grid-template-columns: repeat(2, 1fr); }}
      .metric-card {{ border-radius: 6px; padding: 4mm 5mm; display: flex; flex-direction: column; gap: 3px; }}
      .metric-card .card-label {{ font-size: 7.5pt; font-weight: 600; color: var(--mid-text); text-transform: uppercase; letter-spacing: 0.4px; }}
      .metric-card .card-value {{ font-size: 16pt; font-weight: 700; color: var(--navy); line-height: 1.1; }}
      .metric-card .card-sub {{ font-size: 7pt; color: var(--mid-text); margin-top: 1px; }}
      .card-blue   {{ background: var(--light-blue); }}
      .card-teal   {{ background: var(--light-teal); }}
      .card-orange {{ background: var(--light-orange); }}
      .card-red    {{ background: var(--red-light); }}
      .card-grey   {{ background: var(--grey-bg); }}
      .card-green  {{ background: var(--green-light); }}
      .card-yellow {{ background: var(--yellow-light); }}
      .data-table {{ width: 100%; border-collapse: collapse; font-size: 8.5pt; margin-bottom: 2mm; }}
      .data-table thead tr {{ background: var(--navy); color: #fff; }}
      .data-table thead th {{ padding: 3mm 4mm; text-align: left; font-weight: 700; font-size: 8pt; letter-spacing: 0.3px; border-bottom: 2px solid var(--teal); }}
      .data-table tbody tr:nth-child(even) {{ background: var(--grey-bg); }}
      .data-table tbody tr:nth-child(odd) {{ background: #fff; }}
      .data-table tbody td {{ padding: 2.5mm 4mm; border-bottom: 0.5px solid var(--grey-line); color: var(--dark-text); vertical-align: middle; }}
      .badge {{
        display: inline-block; padding: 1px 7px; border-radius: 20px;
        font-size: 7pt; font-weight: 700; text-transform: uppercase; letter-spacing: 0.4px;
      }}
      .badge-active    {{ background: var(--green-light);  color: var(--green); }}
      .badge-pending   {{ background: var(--yellow-light); color: #92400e; }}
      .badge-inactive  {{ background: var(--grey-bg);      color: var(--mid-text); }}
      .badge-rejected  {{ background: var(--red-light);    color: #9f1239; }}
      .badge-completed {{ background: var(--light-teal);   color: var(--teal); }}
      .badge-expired   {{ background: var(--red-light);    color: #9f1239; }}
      .badge-cancelled {{ background: var(--grey-bg);      color: var(--mid-text); }}
      .recommendation-box {{
        border-left: 4px solid var(--teal); background: var(--light-teal);
        border-radius: 0 6px 6px 0; padding: 3.5mm 5mm; margin-bottom: 4mm;
      }}
      .recommendation-box .rec-label {{ font-size: 7.5pt; font-weight: 700; color: var(--teal); text-transform: uppercase; letter-spacing: 0.5px; margin-bottom: 2px; }}
      .recommendation-box .rec-text {{ font-size: 9.5pt; font-weight: 600; color: var(--dark-text); }}
      .report-footer {{ position: running(reportFooter); background: var(--white); border-top: 1px solid var(--grey-line); padding: 3mm 18mm 2mm 18mm; }}
      .footer-main {{ display: flex; align-items: flex-start; justify-content: space-between; gap: 10mm; padding-bottom: 2mm; border-bottom: 0.5px solid var(--grey-line); }}
      .footer-logo-block {{ display: flex; align-items: center; gap: 6px; }}
      .footer-logo {{ height: 10mm; width: auto; object-fit: contain; }}
      .footer-company {{ font-size: 7.5pt; line-height: 1.6; }}
      .footer-company .company-name {{ font-weight: 700; color: var(--dark-text); font-size: 8.5pt; }}
      .footer-company .powered-by {{ color: #1565C0; font-weight: 700; font-size: 7.5pt; }}
      .footer-company .company-detail {{ color: var(--mid-text); font-size: 7pt; }}
      .footer-bottom {{ display: flex; justify-content: space-between; padding-top: 1.5mm; font-size: 7pt; color: var(--mid-text); }}
      .mt-2 {{ margin-top: 2mm; }}
      .mb-4 {{ margin-bottom: 4mm; }}
      @page {{
        size: A4 portrait; margin: 0 0 34mm 0;
        @bottom-center {{
          content: element(reportFooter);
          width: 100%;
          height: 34mm;
          margin: 0;
          padding: 0;
        }}
      }}
      tr {{ page-break-inside: avoid; }}
    </style>
    </head>
    <body>
      <header class="report-header">
        <div class="header-logo-block">
          <img src="{isf_logo}" class="header-logo" onerror="this.style.display='none'"/>
          <div class="header-brand">
            <div class="brand-name">IndusServiceFlow</div>
            <div class="brand-tagline">Flowing Services, Building Trust</div>
          </div>
        </div>
        <div class="header-center">
          <div class="report-type">REPORT</div>
          <div class="report-subtitle">{html.escape(report_subtitle)}</div>
        </div>
        <div class="header-meta">
          <strong>IndusServiceFlow</strong>
          {generated_at}
        </div>
      </header>
      <main class="page-content">
        <div class="report-title-block">
          <h1>{html.escape(report_title)}</h1>
          <div class="report-meta-row">
            <span>Total Records: <strong>{total_count}</strong></span>
            <span>Generated By: <strong>{html.escape(str(generated_by_username))}</strong></span>
            <span>Generated: <strong>{generated_at}</strong></span>
          </div>
        </div>

        <div class="section">
          <div class="section-title"><span class="section-num">1</span> Summary</div>
          {cards_html}
        </div>

        <div class="section">
          <div class="section-title"><span class="section-num">2</span> {html.escape(table_section_title)}</div>
          {highlight_html}
          <table class="data-table">
            <thead><tr>{header_cells}</tr></thead>
            <tbody>{body_rows}</tbody>
          </table>
        </div>
      </main>
      <footer class="report-footer">
        <div class="footer-main">
          <div class="footer-logo-block">
            <img src="{orchasp_logo}" class="footer-logo" onerror="this.style.display='none'"/>
            <div class="footer-company">
              <div class="company-name">IndusServiceFlow</div>
              <div class="powered-by">Powered by Orchasp Limited</div>
              <div class="company-detail">CIN: L72200TG1994PLC017485</div>
              <div class="company-detail">19 &amp; 20, Moti Valley, Trimulgherry, Secunderabad – 500 015, Telangana, INDIA</div>
              <div class="company-detail">Email: info@orchasp.com &nbsp;|&nbsp; Tel: +91-40-4776 6123 / 124</div>
            </div>
          </div>
        </div>
        <div class="footer-bottom">
          <span>Printed On: {printed_on}</span>
          <span>Confidential – For Internal Use Only</span>
        </div>
      </footer>
    </body>
    </html>
    """

    try:
        from weasyprint import HTML

        pdf_bytes = HTML(string=html_string).write_pdf()
    except ImportError:
        return Response(
            {
                "success": False,
                "message": "PDF export is unavailable because WeasyPrint is not installed.",
            },
            status=status.HTTP_503_SERVICE_UNAVAILABLE,
        )
    except Exception as exc:

        import traceback

        traceback.print_exc()
        return Response(
            {
                "success": False,
                "message": "PDF generation failed.",
                "error_type": type(exc).__name__,
                "error_detail": str(exc),
            },
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


logger = logging.getLogger(__name__)


_STATUS_MAP = {
    "Pending Payment": "pending_payment",
    "Pending Activation": "pending_activation",
    "Active": "active",
    "Expiring Soon": "expiring_soon",
    "Expired": "expired",
    "Locked": "locked",
    "Cancelled": "cancelled",
}
_PAYMENT_STATUS_MAP = {
    "Pending": "pending",
    "Paid": "paid",
    "Failed": "failed",
    "Refunded": "refunded",
}
_BILLING_CYCLE_MAP = {"Monthly": "monthly", "Annual": "annual", "Free Trial": "trial"}
_ORG_STATUS_MAP = {
    "Pending": "pending",
    "Approved": "active",
    "Active": "active",
    "Rejected": "rejected",
    "Inactive": "inactive",
}


@api_view(["GET"])
@permission_classes([AllowAny])
def org_dashboard(request):

    org_id = request.query_params.get("org_id")

    if not org_id:
        return Response(
            {
                "success": False,
                "message": "org_id is required",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    from .serializers import (
        DashboardStatsSerializer,
        QueueLengthTrendSerializer,
        ServiceDistributionSerializer,
        WaitTimeTrendSerializer,
        EmployeeUtilizationSerializer,
        PeakHourSerializer,
        CustomerFlowSerializer,
        WaitDistributionSerializer,
        QueueStatusSerializer,
    )

    return Response(
        {
            "success": True,
            "data": {
                "stats": DashboardStatsSerializer(
                    DashboardService.get_stats(org_id)
                ).data,
                "queue_length_trend": QueueLengthTrendSerializer(
                    DashboardService.get_queue_length_trend(org_id),
                    many=True,
                ).data,
                "service_distribution": ServiceDistributionSerializer(
                    DashboardService.get_service_distribution(org_id),
                    many=True,
                ).data,
                "wait_time_trend": WaitTimeTrendSerializer(
                    DashboardService.get_wait_time_trend(org_id),
                    many=True,
                ).data,
                "employee_utilization": EmployeeUtilizationSerializer(
                    DashboardService.get_employee_utilization(org_id),
                    many=True,
                ).data,
                "peak_hours": PeakHourSerializer(
                    DashboardService.get_peak_hours(org_id),
                    many=True,
                ).data,
                "customer_flow": CustomerFlowSerializer(
                    DashboardService.get_customer_flow(org_id),
                    many=True,
                ).data,
                "wait_distribution": WaitDistributionSerializer(
                    DashboardService.get_wait_distribution(org_id),
                    many=True,
                ).data,
                "queue_status": QueueStatusSerializer(
                    DashboardService.get_queue_status(org_id),
                    many=True,
                ).data,
            },
        }
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def get_my_subscription_status(request):
    organization = getattr(request.user, "organization", None)
    if organization is None:
        return Response(
            {"success": False, "message": "Organization is required."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    subscription = (
        organization.subscriptions.select_related("plan")
        .order_by("-created_on")
        .first()
    )
    if subscription is None:
        return Response(
            {
                "success": True,
                "data": {
                    "organization_status": _ORG_STATUS_MAP.get(
                        organization.status, organization.status.lower()
                    ),
                    "subscription_status": "pending_payment",
                    "payment_status": "pending",
                    "plan_name": "No Plan",
                    "billing_cycle": "trial",
                    "employee_limit": None,
                    "queue_limit": None,
                    "employee_count": organization.employees.filter(
                        status="Active"
                    ).count(),
                    "queue_count": 0,
                    "features": [],
                    "current_period_end": None,
                    "grace_period_end": None,
                    "days_remaining": None,
                },
            },
            status=status.HTTP_200_OK,
        )

    return Response(
        {
            "success": True,
            "data": {
                "organization_status": _ORG_STATUS_MAP.get(
                    organization.status, organization.status.lower()
                ),
                "subscription_status": _STATUS_MAP.get(
                    subscription.status, subscription.status.lower()
                ),
                "payment_status": _PAYMENT_STATUS_MAP.get(
                    subscription.payment_status, subscription.payment_status.lower()
                ),
                "plan_name": subscription.plan.plan_name,
                "billing_cycle": _BILLING_CYCLE_MAP.get(
                    subscription.billing_cycle, subscription.billing_cycle.lower()
                ),
                "employee_limit": subscription.plan.employee_limit,
                "queue_limit": subscription.plan.queue_limit,
                "employee_count": organization.employees.filter(
                    status="Active"
                ).count(),
                "queue_count": 0,
                "features": subscription.plan.features or [],
                "current_period_end": subscription.next_payment_date,
                "grace_period_end": None,
                "days_remaining": None,
            },
        },
        status=status.HTTP_200_OK,
    )


@api_view(["POST"])
@permission_classes([AllowAny])
def register_organization(request):

    plan_id = request.data.get("plan_id")
    billing_cycle = request.data.get("billing_cycle")

    if not plan_id or billing_cycle not in ("monthly", "annual", "trial"):
        return Response(
            {
                "success": False,
                "message": "plan_id and a valid billing_cycle are required.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    plan = Plan.objects.filter(id=plan_id, status="Active").first()
    if not plan:
        return Response(
            {"success": False, "message": "Selected plan is not available."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if billing_cycle == "trial" and not plan.trial_days:
        return Response(
            {"success": False, "message": "This plan does not offer a free trial."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    serializer = OrganizationRegistrationSerializer(data=request.data)

    if not serializer.is_valid():
        return Response(
            {"success": False, "errors": serializer.errors},
            status=status.HTTP_400_BAD_REQUEST,
        )

    result = serializer.save()
    organization = result["organization"]
    contact_person_name = result["contact_person_name"]
    contact_email = result["contact_email"]

    is_trial = billing_cycle == "trial"
    amount = (
        0
        if is_trial
        # plan.annual_total (monthly_price x 12) rather than the raw
        # plan.annual_price column - see Plan.annual_total for why: the
        # stored column could be configured independently of monthly
        # price and drift from what the customer is actually shown/billed.
        else (plan.monthly_price if billing_cycle == "monthly" else plan.annual_total)
    )

    db_billing_cycle = {
        "monthly": "Monthly",
        "annual": "Annual",
        "trial": "Free Trial",
    }[billing_cycle]

    subscription = Subscription.objects.create(
        organization=organization,
        plan=plan,
        billing_cycle=db_billing_cycle,
        amount=amount,
        is_free_trial=is_trial,
        payment_status="Paid" if is_trial else "Pending",
        status="Pending Activation" if is_trial else "Pending Payment",
        start_date=date.today(),
        next_payment_date=(
            date.today() + timedelta(days=plan.trial_days)
            if is_trial
            else date.today() + timedelta(days=365 if billing_cycle == "annual" else 30)
        ),
        created_by="System",
    )

    razorpay_payload = None

    if not is_trial:
        fake_order_id = f"order_fake_{subscription.id}"

        qr_text = f"Payment for {organization.organization_name} - {plan.plan_name} - Rs.{amount}"
        qr_url = (
            "https://api.qrserver.com/v1/create-qr-code/"
            f"?size=300x300&data={qr_text.replace(' ', '+')}"
        )

        subscription.razorpay_order_id = fake_order_id
        subscription.razorpay_qr_code_url = qr_url
        subscription.save()

        razorpay_payload = {
            "order_id": fake_order_id,
            "qr_code_url": qr_url,
            "amount": float(amount),
            "currency": "INR",
        }

    subject = "Organization Registration Successful"

    message = f"""
Hello {contact_person_name},

Thank you for registering your organization with IndusServiceFlow.

Your organization registration request has been submitted successfully.

Organization Name:
{organization.organization_name}

Current Status:
Pending Approval

Please wait until the Super Admin approves your organization.

Once approved, your Username and Password will be sent to this email.

Regards,
IndusServiceFlow Team
"""

    try:
        send_mail(
            subject,
            message,
            settings.DEFAULT_FROM_EMAIL,
            [contact_email],
            fail_silently=False,
        )

    except Exception:
        logger.exception(
            "Failed to send registration confirmation email " "for org_id=%s",
            organization.id,
        )

    notify_super_admin(
        organization,
        title="New Organization Registration",
        message=(
            f"{organization.organization_name} has registered "
            f"and is pending approval."
        ),
        notification_type="Organization Registration",
        actor=request.user,
    )

    return Response(
        {
            "success": True,
            "message": "Organization registered successfully.",
            "organization": {
                "id": organization.id,
                "organization_name": organization.organization_name,
                "status": _ORG_STATUS_MAP.get(organization.status, "pending"),
            },
            "subscription": {
                "id": subscription.id,
                "status": _STATUS_MAP[subscription.status],
                "payment_status": _PAYMENT_STATUS_MAP[subscription.payment_status],
                "is_free_trial": subscription.is_free_trial,
                "plan_name": plan.plan_name,
                "billing_cycle": _BILLING_CYCLE_MAP[subscription.billing_cycle],
                "amount": float(subscription.amount),
            },
            "razorpay": razorpay_payload,
        },
        status=status.HTTP_201_CREATED,
    )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def renew_subscription(request):

    org_id = getattr(request.user, "organization_id", None)

    if org_id is None:
        return Response(
            {
                "success": False,
                "message": "This account is not linked to an organization.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    organization = Organization.objects.filter(id=org_id).first()

    if organization is None:
        return Response(
            {"success": False, "message": "Organization not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    plan_id = request.data.get("plan_id")
    billing_cycle = request.data.get("billing_cycle")

    if not plan_id or billing_cycle not in ("monthly", "annual", "trial"):
        return Response(
            {
                "success": False,
                "message": "plan_id and a valid billing_cycle are required.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    plan = Plan.objects.filter(id=plan_id, status="Active").first()
    if not plan:
        return Response(
            {"success": False, "message": "Selected plan is not available."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if billing_cycle == "trial" and not plan.trial_days:
        return Response(
            {"success": False, "message": "This plan does not offer a free trial."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    is_trial = billing_cycle == "trial"
    amount = (
        0
        if is_trial
        # plan.annual_total (monthly_price x 12) rather than the raw
        # plan.annual_price column - see Plan.annual_total for why: the
        # stored column could be configured independently of monthly
        # price and drift from what the customer is actually shown/billed.
        else (plan.monthly_price if billing_cycle == "monthly" else plan.annual_total)
    )

    db_billing_cycle = {
        "monthly": "Monthly",
        "annual": "Annual",
        "trial": "Free Trial",
    }[billing_cycle]

    subscription = Subscription.objects.create(
        organization=organization,
        plan=plan,
        billing_cycle=db_billing_cycle,
        amount=amount,
        is_free_trial=is_trial,
        payment_status="Paid" if is_trial else "Pending",
        status="Pending Activation" if is_trial else "Pending Payment",
        start_date=date.today(),
        next_payment_date=(
            date.today() + timedelta(days=plan.trial_days)
            if is_trial
            else date.today() + timedelta(days=365 if billing_cycle == "annual" else 30)
        ),
        created_by=request.user.username,
    )

    razorpay_payload = None

    if not is_trial:

        fake_order_id = f"order_fake_{subscription.id}"

        qr_text = f"Renewal for {organization.organization_name} - {plan.plan_name} - Rs.{amount}"
        qr_url = (
            "https://api.qrserver.com/v1/create-qr-code/"
            f"?size=300x300&data={qr_text.replace(' ', '+')}"
        )

        subscription.razorpay_order_id = fake_order_id
        subscription.razorpay_qr_code_url = qr_url
        subscription.save()

        razorpay_payload = {
            "order_id": fake_order_id,
            "qr_code_url": qr_url,
            "amount": float(amount),
            "currency": "INR",
        }

    log_action(
        request.user,
        "Create",
        "Subscriptions",
        target_info=f"{organization.organization_name} - {plan.plan_name} (renewal)",
    )

    notify_super_admin(
        organization,
        title="Subscription Renewal",
        message=(
            f"{organization.organization_name} has started a plan renewal "
            f"({plan.plan_name}, {db_billing_cycle})."
        ),
        notification_type="Subscription Renewal",
        actor=request.user,
    )

    return Response(
        {
            "success": True,
            "message": "Renewal started.",
            "data": {
                "organization": {
                    "id": organization.id,
                    "organization_name": organization.organization_name,
                    "status": _ORG_STATUS_MAP.get(organization.status, "pending"),
                },
                "subscription": {
                    "id": subscription.id,
                    "status": _STATUS_MAP[subscription.status],
                    "payment_status": _PAYMENT_STATUS_MAP[subscription.payment_status],
                    "is_free_trial": subscription.is_free_trial,
                    "plan_name": plan.plan_name,
                    "billing_cycle": _BILLING_CYCLE_MAP[subscription.billing_cycle],
                    "amount": float(subscription.amount),
                },
                "razorpay": razorpay_payload,
            },
        },
        status=status.HTTP_201_CREATED,
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def get_renewal_payment_status(request, subscription_id):

    org_id = getattr(request.user, "organization_id", None)

    subscription = Subscription.objects.filter(
        id=subscription_id,
        organization_id=org_id,
    ).first()

    if not subscription:
        return Response(
            {"success": False, "message": "Subscription not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    return Response(
        {
            "success": True,
            "data": {
                "payment_status": _PAYMENT_STATUS_MAP[subscription.payment_status],
            },
        },
        status=status.HTTP_200_OK,
    )


def _build_subscription_receipt_pdf(subscription):
    """
    Small standalone PDF receipt for a Paid subscription (registration OR
    renewal/upgrade — both live in the Subscription table). Deliberately
    separate from _build_list_report_pdf above, which is built for tabular
    admin reports (KPI cards + data table), not a single-record receipt.
    """
    static_dir = os.path.join(django_settings.BASE_DIR, "statics")
    isf_logo = "file:///" + pathname2url(
        os.path.join(static_dir, "indusserviseflow logo.png")
    ).lstrip("/")

    organization = subscription.organization
    receipt_no = f"SUB-{str(subscription.id).zfill(5)}"
    issued_on = (subscription.updated_on or subscription.created_on).strftime(
        "%d %b %Y, %I:%M %p"
    )
    billing_cycle_label = _BILLING_CYCLE_MAP.get(
        subscription.billing_cycle, subscription.billing_cycle
    )

    html_string = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="UTF-8"/>
    <style>
      * {{ box-sizing: border-box; margin: 0; padding: 0; }}
      body {{ font-family:'Helvetica Neue',Helvetica,Arial,sans-serif; color:#0F172A; }}
      .page {{ padding: 14mm 18mm; }}
      .header {{ display:flex; align-items:center; justify-content:space-between; border-bottom:2.5px solid #0d9488; padding-bottom:6mm; margin-bottom:8mm; }}
      .brand {{ display:flex; align-items:center; gap:10px; }}
      .brand img {{ height:14mm; width:auto; object-fit:contain; }}
      .brand-name {{ font-size:14pt; font-weight:700; color:#1B2A6B; }}
      .doc-title {{ text-align:right; }}
      .doc-title h1 {{ font-size:16pt; color:#1B2A6B; }}
      .doc-title .meta {{ font-size:8.5pt; color:#475569; margin-top:2px; }}
      .row {{ display:flex; justify-content:space-between; margin-bottom:4mm; font-size:9.5pt; }}
      .row .label {{ color:#475569; }}
      .row .value {{ font-weight:600; }}
      table {{ width:100%; border-collapse:collapse; margin-top:6mm; font-size:9.5pt; }}
      th {{ background:#1B2A6B; color:#fff; text-align:left; padding:3mm 4mm; }}
      td {{ padding:3mm 4mm; border-bottom:0.5px solid #CBD5E1; }}
      .total-row td {{ font-weight:700; font-size:11pt; border-top:2px solid #0d9488; border-bottom:none; }}
      .paid-badge {{ display:inline-block; padding:1.5px 9px; border-radius:20px; background:#DCFCE7; color:#16a34a; font-size:8pt; font-weight:700; text-transform:uppercase; }}
      .footer {{ margin-top:12mm; font-size:7.5pt; color:#94A3B8; text-align:center; }}
    </style>
    </head>
    <body>
      <div class="page">
        <div class="header">
          <div class="brand">
            <img src="{isf_logo}" onerror="this.style.display='none'"/>
            <div class="brand-name">IndusServiceFlow</div>
          </div>
          <div class="doc-title">
            <h1>Payment Receipt</h1>
            <div class="meta">Receipt No: {html.escape(receipt_no)}</div>
            <div class="meta">Issued: {html.escape(issued_on)}</div>
          </div>
        </div>

        <div class="row"><span class="label">Billed To</span><span class="value">{html.escape(organization.organization_name)}</span></div>
        <div class="row"><span class="label">Plan</span><span class="value">{html.escape(subscription.plan.plan_name)}</span></div>
        <div class="row"><span class="label">Billing Cycle</span><span class="value">{html.escape(str(billing_cycle_label).capitalize())}</span></div>
        <div class="row"><span class="label">Payment Status</span><span class="value"><span class="paid-badge">Paid</span></span></div>

        <table>
          <thead><tr><th>Description</th><th>Amount</th></tr></thead>
          <tbody>
            <tr><td>{html.escape(subscription.plan.plan_name)} — {html.escape(str(billing_cycle_label).capitalize())} subscription</td><td>&#8377; {subscription.amount}</td></tr>
            <tr class="total-row"><td>Total Paid</td><td>&#8377; {subscription.amount}</td></tr>
          </tbody>
        </table>

        <div class="footer">
          IndusServiceFlow &middot; Powered by Orchasp Limited &middot; This is a system-generated receipt.
        </div>
      </div>
    </body>
    </html>
    """

    from weasyprint import HTML

    return HTML(string=html_string).write_pdf()


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def download_subscription_receipt(request, subscription_id):

    org_id = getattr(request.user, "organization_id", None)

    subscription = (
        Subscription.objects.filter(
            id=subscription_id,
            organization_id=org_id,
        )
        .select_related("plan", "organization")
        .first()
    )

    if not subscription:
        return Response(
            {"success": False, "message": "Subscription not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    if subscription.payment_status != "Paid":
        return Response(
            {
                "success": False,
                "message": "No receipt available yet — payment hasn't been confirmed for this subscription.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        pdf_bytes = _build_subscription_receipt_pdf(subscription)
    except ImportError:
        return Response(
            {
                "success": False,
                "message": "Receipt generation is unavailable because WeasyPrint is not installed.",
            },
            status=status.HTTP_503_SERVICE_UNAVAILABLE,
        )
    except Exception as exc:
        import traceback

        traceback.print_exc()
        return Response(
            {
                "success": False,
                "message": "Receipt generation failed.",
                "error_type": type(exc).__name__,
                "error_detail": str(exc),
            },
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="receipt_SUB-{str(subscription.id).zfill(5)}.pdf"'
    )
    return response


@api_view(["GET"])
@permission_classes([AllowAny])
def download_public_subscription_receipt(request, subscription_id):

    subscription = (
        Subscription.objects.filter(
            id=subscription_id,
        )
        .select_related("plan", "organization")
        .first()
    )

    if not subscription:
        return Response(
            {"success": False, "message": "Subscription not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    if subscription.payment_status != "Paid":
        return Response(
            {
                "success": False,
                "message": "No receipt available yet — payment hasn't been confirmed for this subscription.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        pdf_bytes = _build_subscription_receipt_pdf(subscription)
    except ImportError:
        return Response(
            {
                "success": False,
                "message": "Receipt generation is unavailable because WeasyPrint is not installed.",
            },
            status=status.HTTP_503_SERVICE_UNAVAILABLE,
        )
    except Exception as exc:
        import traceback

        traceback.print_exc()
        return Response(
            {
                "success": False,
                "message": "Receipt generation failed.",
                "error_type": type(exc).__name__,
                "error_detail": str(exc),
            },
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="receipt_SUB-{str(subscription.id).zfill(5)}.pdf"'
    )
    return response


GRACE_PERIOD_DAYS = 7


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def get_my_subscription_status(request):

    org_id = getattr(request.user, "organization_id", None)

    if org_id is None:
        return Response(
            {
                "success": False,
                "message": "This account is not linked to an organization.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    organization = Organization.objects.filter(id=org_id).first()

    if organization is None:
        return Response(
            {"success": False, "message": "Organization not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    subscription = (
        Subscription.objects.filter(organization=organization, status="Active")
        .select_related("plan")
        .order_by("-created_on")
        .first()
    )
    if subscription is None:
        subscription = (
            Subscription.objects.filter(organization=organization)
            .select_related("plan")
            .order_by("-created_on")
            .first()
        )

    if subscription is None:
        return Response(
            {
                "success": False,
                "message": "No subscription found for this organization.",
            },
            status=status.HTTP_404_NOT_FOUND,
        )

    plan = subscription.plan

    from services.models import Service

    queue_count = Service.objects.filter(
        organization=organization,
        status="Active",
    ).count()

    sub_status = _STATUS_MAP.get(subscription.status, "pending_payment")

    grace_period_end = None
    days_remaining = None

    if sub_status == "expired" and subscription.next_payment_date:
        grace_period_end = subscription.next_payment_date + timedelta(
            days=GRACE_PERIOD_DAYS
        )
        days_remaining = max(0, (grace_period_end - date.today()).days)
    elif sub_status == "expiring_soon" and subscription.next_payment_date:
        days_remaining = max(0, (subscription.next_payment_date - date.today()).days)

    return Response(
        {
            "success": True,
            "data": {
                "organization_status": _ORG_STATUS_MAP.get(
                    organization.status, "pending"
                ),
                "subscription_status": sub_status,
                "payment_status": _PAYMENT_STATUS_MAP.get(
                    subscription.payment_status, "pending"
                ),
                "plan_name": plan.plan_name,
                "billing_cycle": _BILLING_CYCLE_MAP.get(
                    subscription.billing_cycle, "monthly"
                ),
                "employee_limit": plan.employee_limit,
                "queue_limit": plan.queue_limit,
                "employee_count": subscription.live_employee_count(),
                "queue_count": queue_count,
                "features": plan.features,
                "current_period_end": subscription.next_payment_date,
                "grace_period_end": grace_period_end,
                "days_remaining": days_remaining,
            },
        },
        status=status.HTTP_200_OK,
    )


@api_view(["GET"])
@permission_classes([AllowAny])
def registration_payment_status(request, subscription_id):
    subscription = Subscription.objects.filter(id=subscription_id).first()

    if not subscription:
        return Response(
            {"success": False, "message": "Subscription not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    return Response(
        {
            "success": True,
            "payment_status": _PAYMENT_STATUS_MAP[subscription.payment_status],
        },
        status=status.HTTP_200_OK,
    )


@api_view(["POST"])
@permission_classes([AllowAny])
def simulate_registration_payment(request, subscription_id):
    subscription = Subscription.objects.filter(id=subscription_id).first()

    if not subscription:
        return Response(
            {"success": False, "message": "Subscription not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    if subscription.is_free_trial:
        return Response(
            {
                "success": False,
                "message": "Free trial subscriptions don't require payment.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    if subscription.payment_status != "Paid":
        subscription.payment_status = "Paid"

        organization = subscription.organization
        is_renewal_or_upgrade = (
            organization is not None and organization.status == "Active"
        )

        if is_renewal_or_upgrade:
            subscription.status = "Active"

            if subscription.billing_cycle == "Monthly":
                subscription.monthly_revenue = subscription.amount
            elif subscription.billing_cycle == "Annual":
                subscription.monthly_revenue = subscription.amount / 12
            else:
                subscription.monthly_revenue = 0

            Subscription.objects.filter(
                organization=organization,
                status="Active",
            ).exclude(id=subscription.id).update(status="Expired")

            subscription.save()

            notify_org_admins(
                organization=organization,
                title="Subscription Activated",
                message=(
                    f"Your renewed subscription "
                    f"({subscription.plan.plan_name}, {subscription.billing_cycle}) "
                    f"is now active."
                ),
                notification_type="Organization Status Update",
            )
        else:
            subscription.status = "Pending Activation"
            subscription.save()

    return Response(
        {
            "success": True,
            "payment_status": _PAYMENT_STATUS_MAP[subscription.payment_status],
            "subscription_status": _STATUS_MAP[subscription.status],
        },
        status=status.HTTP_200_OK,
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def get_organization(request, pk):

    organization = get_object_or_404(
        _with_latest_subscription(Organization.objects),
        pk=pk,
        is_deleted=False,
    )

    serializer = OrganizationSerializer(organization)

    return Response(
        {
            "success": True,
            "data": serializer.data,
        },
        status=status.HTTP_200_OK,
    )


@api_view(["GET"])
@permission_classes([AllowAny])
def get_organizations(request):

    organizations = _with_latest_subscription(
        Organization.objects.filter(is_deleted=False)
    ).order_by("-id")

    search = request.query_params.get("search")

    status_filter = request.query_params.get("status")
    category = request.query_params.get("category")
    category_name = request.query_params.get("category_name")
    org_type = request.query_params.get("type")

    if not request.user or not request.user.is_authenticated:

        status_filter = "Active"
        org_type = None

    if search:

        organizations = organizations.filter(
            Q(organization_name__icontains=search)
            | Q(email__icontains=search)
            | Q(city__icontains=search)
        )

    if status_filter:

        organizations = organizations.filter(status=status_filter)

    if category:

        organizations = organizations.filter(category_id=category)

    if category_name:

        organizations = organizations.filter(
            category__category_name__iexact=category_name
        )

    if org_type:

        if org_type.lower() == "request":

            organizations = organizations.filter(status="Pending")

        elif org_type.lower() == "organization":

            organizations = organizations.exclude(status="Pending")

    organizations = organizations.order_by("-id")

    all_orgs = Organization.objects.filter(is_deleted=False)
    counts = {
        "total": all_orgs.count(),
        "pending": all_orgs.filter(status__iexact="Pending").count(),
        "active": all_orgs.filter(status__iexact="Active").count(),
        "rejected": all_orgs.filter(status__iexact="Rejected").count(),
    }

    page_param = request.query_params.get("page")

    if page_param is not None:

        try:
            page = int(page_param)

        except (TypeError, ValueError):
            page = 1

        page = max(page, 1)

        try:
            page_size = int(
                request.query_params.get(
                    "page_size",
                    10,
                )
            )

        except (TypeError, ValueError):
            page_size = 10

        page_size = max(
            1,
            min(page_size, 100),
        )

        paginator = Paginator(
            organizations,
            page_size,
        )

        page_obj = paginator.get_page(page)

        serializer = OrganizationSerializer(
            page_obj.object_list,
            many=True,
        )

        return Response(
            {
                "success": True,
                "message": (
                    "Organizations retrieved successfully."
                    if paginator.count
                    else "No organizations found."
                ),
                "pagination": {
                    "current_page": page_obj.number,
                    "total_pages": paginator.num_pages,
                    "total_records": paginator.count,
                    "page_size": page_size,
                    "has_next": page_obj.has_next(),
                    "has_previous": page_obj.has_previous(),
                },
                "counts": counts,
                "data": serializer.data,
            },
            status=status.HTTP_200_OK,
        )

    serializer = OrganizationSerializer(
        organizations,
        many=True,
    )

    return Response(
        {
            "success": True,
            "message": (
                "Organizations retrieved successfully."
                if organizations.exists()
                else "No organizations found."
            ),
            "count": organizations.count(),
            "data": serializer.data,
        },
        status=status.HTTP_200_OK,
    )


@api_view(["PUT"])
@permission_classes([IsAuthenticated])
def approve_organization(request, pk):

    if request.user.role != "SUPER_ADMIN":

        return Response(
            {
                "success": False,
                "message": ("Only Super Admin can approve organizations."),
            },
            status=status.HTTP_403_FORBIDDEN,
        )

    organization = get_object_or_404(
        Organization,
        pk=pk,
        is_deleted=False,
    )

    if organization.status == "Active":

        return Response(
            {
                "success": False,
                "message": ("This organization is already approved."),
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    # Look this up *before* changing anything. Previously the organization
    # (and its subscription) were flipped to Active first and only then was
    # the ORG_ADMIN looked up — if it was missing, this returned a 404 but
    # left the organization sitting at "Active" with no admin able to log in
    # or receive credentials, and no subscription visibly "Pending" either.
    # That's exactly how an org silently drops out of the Users list while
    # still counting as an active organization elsewhere. Checking first
    # means a missing admin blocks the approval instead of half-completing it.
    org_admin = User.objects.filter(
        organization=organization,
        role="ORG_ADMIN",
    ).first()

    if not org_admin:

        return Response(
            {
                "success": False,
                "message": (
                    "Organization Admin not found. This organization can't "
                    "be approved until it has one — run "
                    "'python manage.py backfill_missing_admins' to create a "
                    "placeholder admin for it, then try again."
                ),
            },
            status=status.HTTP_404_NOT_FOUND,
        )

    with transaction.atomic():

        organization.status = "Active"
        organization.updated_by = request.user.username
        organization.save()

        subscription = (
            Subscription.objects.filter(organization=organization)
            .order_by("-id")
            .first()
        )

        if subscription and subscription.status != "Active":

            subscription.status = "Active"

            if subscription.billing_cycle == "Monthly":
                subscription.monthly_revenue = subscription.amount
            elif subscription.billing_cycle == "Annual":
                subscription.monthly_revenue = subscription.amount / 12
            else:
                subscription.monthly_revenue = 0

            if subscription.payment_status != "Paid":
                subscription.payment_status = "Paid"

            subscription.updated_by = request.user.username
            subscription.save()

    log_action(
        request.user,
        "Approve",
        "Organizations",
        target_info=organization.organization_name,
    )

    first_name = org_admin.name.split()[0].lower() if org_admin.name else "admin"

    while True:

        username = f"{first_name}{random.randint(1000, 9999)}"

        if not User.objects.filter(username=username).exclude(id=org_admin.id).exists():

            break

    password = "".join(
        random.choices(
            string.ascii_letters + string.digits,
            k=10,
        )
    )

    org_admin.username = username
    org_admin.set_password(password)
    org_admin.status = "Active"
    org_admin.updated_by = request.user.username
    org_admin.save()

    notify_org_admins(
        organization=organization,
        title="Organization Approved",
        message=(
            f"Your organization "
            f"{organization.organization_name} has been approved. "
            f"You can now access the organization portal."
        ),
        notification_type="Organization Status Update",
        actor=request.user,
    )

    subject = "Organization Approved - Login Credentials"

    message = f"""
Hello {org_admin.name},

Congratulations!

Your organization has been approved successfully.

You can now login to IndusServiceFlow.

Username:
{username}

Password:
{password}

Please change your password after your first login.

Regards,
IndusServiceFlow Team
"""

    email_status = "sent"

    try:

        send_mail(
            subject,
            message,
            settings.DEFAULT_FROM_EMAIL,
            [org_admin.email],
            fail_silently=False,
        )

    except Exception:

        email_status = "failed"

        logger.exception(
            "Failed to send approval credentials email " "for organization_id=%s",
            organization.id,
        )

    response_data = {
        "success": True,
        "message": (
            "Organization approved successfully."
            if email_status == "sent"
            else (
                "Organization approved successfully, but the "
                "credentials email could not be sent. Share the "
                "login details below with the organization admin "
                "manually - they won't be shown again."
            )
        ),
        "email_status": email_status,
    }

    if email_status == "failed":

        response_data["credentials"] = {
            "username": username,
            "password": password,
        }

    return Response(
        response_data,
        status=status.HTTP_200_OK,
    )


@api_view(["PUT"])
@permission_classes([IsAuthenticated])
def reject_organization(request, pk):

    if request.user.role != "SUPER_ADMIN":

        return Response(
            {
                "success": False,
                "message": ("Only Super Admin can reject organizations."),
            },
            status=status.HTTP_403_FORBIDDEN,
        )

    organization = get_object_or_404(
        Organization,
        pk=pk,
        is_deleted=False,
    )

    with transaction.atomic():

        organization.status = "Rejected"
        organization.updated_by = request.user.username
        organization.save()

        subscription = (
            Subscription.objects.filter(organization=organization)
            .order_by("-id")
            .first()
        )

        refund_due = bool(subscription and subscription.payment_status == "Paid")

        if subscription and (subscription.status != "Cancelled" or refund_due):
            subscription.status = "Cancelled"
            if refund_due:
                subscription.payment_status = "Refunded"
            subscription.updated_by = request.user.username
            subscription.save(
                update_fields=["status", "payment_status", "updated_by", "updated_on"]
            )

    log_action(
        request.user,
        "Reject",
        "Organizations",
        target_info=organization.organization_name,
    )

    org_admin = User.objects.filter(
        organization=organization,
        role="ORG_ADMIN",
    ).first()

    # A rejection is still valid even if this org has no linked admin (e.g.
    # it was already missing one — see backfill_missing_admins). The org and
    # its subscription above are already rejected/cancelled regardless, so
    # this no longer hard-fails with a 404 that would misleadingly suggest
    # the rejection itself didn't go through.
    if not org_admin:

        notify_org_admins(
            organization=organization,
            title="Organization Registration Rejected",
            message=(
                f"Your organization "
                f"{organization.organization_name} registration "
                f"was rejected."
                + (
                    " Your payment will be refunded within 24 hours."
                    if refund_due
                    else ""
                )
            ),
            notification_type="Organization Status Update",
            actor=request.user,
        )

        return Response(
            {
                "success": True,
                "message": (
                    "Organization rejected successfully. No linked admin "
                    "user was found, so no rejection email could be sent."
                ),
                "email_status": "not_applicable",
            },
            status=status.HTTP_200_OK,
        )

    org_admin.status = "Inactive"
    org_admin.updated_by = request.user.username
    org_admin.save()

    notify_org_admins(
        organization=organization,
        title="Organization Registration Rejected",
        message=(
            f"Your organization "
            f"{organization.organization_name} registration "
            f"was rejected."
            + (" Your payment will be refunded within 24 hours." if refund_due else "")
        ),
        notification_type="Organization Status Update",
        actor=request.user,
    )

    subject = "Organization Registration Rejected"

    refund_note = (
        f"""
We can confirm your payment for this registration was received. Your money will be refunded within 24 hours to the original payment method.
"""
        if refund_due
        else ""
    )

    message = f"""
Hello {org_admin.name},

Thank you for registering with IndusServiceFlow.

After reviewing your registration request, we regret to inform you that your organization registration has not been approved.

Organization Name:
{organization.organization_name}
{refund_note}
If you believe this was a mistake or would like to submit a new request, please contact the system administrator or register again with the correct details.

Regards,
IndusServiceFlow Team
"""

    email_status = "sent"

    try:

        send_mail(
            subject,
            message,
            settings.DEFAULT_FROM_EMAIL,
            [org_admin.email],
            fail_silently=False,
        )

    except Exception:

        email_status = "failed"

        logger.exception(
            "Failed to send rejection email " "for organization_id=%s",
            organization.id,
        )

    return Response(
        {
            "success": True,
            "message": (
                "Organization rejected successfully."
                if email_status == "sent"
                else (
                    "Organization rejected successfully, but the "
                    "notification email could not be sent to the "
                    "organization admin."
                )
            ),
            "email_status": email_status,
        },
        status=status.HTTP_200_OK,
    )


@api_view(["DELETE"])
@permission_classes([IsAuthenticated])
def delete_organization(request, pk):

    if request.user.role != "SUPER_ADMIN":

        return Response(
            {
                "success": False,
                "message": ("Only Super Admin can delete organizations."),
            },
            status=status.HTTP_403_FORBIDDEN,
        )

    organization = get_object_or_404(
        Organization,
        pk=pk,
        is_deleted=False,
    )

    organization_name = organization.organization_name

    organization.delete()

    log_action(
        request.user,
        "Delete",
        "Organizations",
        target_info=organization_name,
    )

    return Response(
        {
            "success": True,
            "message": ("Organization and its linked users " "deleted successfully."),
        },
        status=status.HTTP_200_OK,
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def export_organizations(request):

    organizations = Organization.objects.filter(is_deleted=False).order_by("id")

    search = request.query_params.get("search")
    status_filter = request.query_params.get("status")
    category = request.query_params.get("category")
    category_name = request.query_params.get("category_name")
    org_type = request.query_params.get("type")

    if search:

        organizations = organizations.filter(
            Q(organization_name__icontains=search)
            | Q(email__icontains=search)
            | Q(city__icontains=search)
        )

    if status_filter:

        organizations = organizations.filter(status=status_filter)

    if category:

        organizations = organizations.filter(
            category__category_name__icontains=category
        )

    if category_name:

        organizations = organizations.filter(
            category__category_name__iexact=category_name
        )

    if org_type:

        if org_type.lower() == "request":

            organizations = organizations.filter(status="Pending")

        elif org_type.lower() == "organization":

            organizations = organizations.exclude(status="Pending")

    export_format = request.query_params.get(
        "export_format",
        "csv",
    ).lower()

    report_title = "ORGANIZATIONS REPORT"
    date_str = timezone.now().strftime("%Y-%m-%d")
    generated_by = f"Generated By: {request.user.username}"

    headers = [
        "ID",
        "Organization Name",
        "Category",
        "Email",
        "Mobile",
        "City",
        "State",
        "Country",
        "Status",
    ]

    if export_format == "csv":

        response = HttpResponse(content_type="text/csv")

        response["Content-Disposition"] = (
            f'attachment; filename="organizations_{date_str}.csv"'
        )

        writer = csv.writer(response)

        # ---- Heading ----
        writer.writerow([report_title])
        writer.writerow([generated_by])
        writer.writerow([])

        writer.writerow(headers)

        for sno, org in enumerate(organizations, 1):

            writer.writerow(
                [
                    sno,
                    org.organization_name,
                    (org.category.category_name if org.category else ""),
                    org.email,
                    org.mobile,
                    org.city,
                    org.state,
                    org.country,
                    org.status,
                ]
            )

        log_action(request.user, "Export", "Organizations")
        return response

    elif export_format == "excel":

        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Organizations"

        last_col = len(headers)

        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        title_cell = sheet.cell(row=1, column=1)
        title_cell.value = report_title
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="center")

        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        sub_cell = sheet.cell(row=2, column=1)
        sub_cell.value = generated_by
        sub_cell.font = Font(italic=True, size=10)
        sub_cell.alignment = Alignment(horizontal="center")

        header_row = 4

        for col_num, header in enumerate(
            headers,
            1,
        ):

            cell = sheet.cell(
                row=header_row,
                column=col_num,
            )

            cell.value = header
            cell.font = Font(bold=True)

        row_num = header_row + 1

        for sno, org in enumerate(organizations, 1):

            sheet.cell(
                row=row_num,
                column=1,
            ).value = sno

            sheet.cell(
                row=row_num,
                column=2,
            ).value = org.organization_name

            sheet.cell(
                row=row_num,
                column=3,
            ).value = (
                org.category.category_name if org.category else ""
            )

            sheet.cell(
                row=row_num,
                column=4,
            ).value = org.email

            sheet.cell(
                row=row_num,
                column=5,
            ).value = org.mobile

            sheet.cell(
                row=row_num,
                column=6,
            ).value = org.city

            sheet.cell(
                row=row_num,
                column=7,
            ).value = org.state

            sheet.cell(
                row=row_num,
                column=8,
            ).value = org.country

            sheet.cell(
                row=row_num,
                column=9,
            ).value = org.status

            row_num += 1

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument." "spreadsheetml.sheet"
            )
        )

        response["Content-Disposition"] = (
            f'attachment; filename="organizations_{date_str}.xlsx"'
        )

        workbook.save(response)

        log_action(request.user, "Export", "Organizations")
        return response

    elif export_format == "pdf":

        organizations_list = list(organizations)

        total_orgs = len(organizations_list)
        active_count = sum(
            1
            for o in organizations_list
            if (o.status or "").strip().lower() == "active"
        )
        pending_count = sum(
            1
            for o in organizations_list
            if (o.status or "").strip().lower() == "pending"
        )
        other_count = total_orgs - active_count - pending_count

        summary_cards = [
            {
                "label": "Total Organizations",
                "value": total_orgs,
                "sub": "All organizations",
                "css": "card-blue",
            },
            {
                "label": "Active",
                "value": active_count,
                "sub": "Currently active",
                "css": "card-green",
            },
            {
                "label": "Pending",
                "value": pending_count,
                "sub": "Awaiting approval",
                "css": "card-yellow",
            },
            {
                "label": "Rejected / Inactive",
                "value": other_count,
                "sub": "Rejected, inactive or approved",
                "css": "card-red",
            },
        ]

        pdf_headers = [
            "ID",
            "Organization",
            "Category",
            "Email",
            "Mobile",
            "City",
            "Status",
        ]

        rows = [
            [
                sno,
                org.organization_name,
                org.category.category_name if org.category else "",
                org.email,
                org.mobile,
                org.city,
                org.status,
            ]
            for sno, org in enumerate(organizations_list, 1)
        ]

        response = _build_list_report_pdf(
            report_title=report_title,
            report_subtitle="Operational Analytics & Performance Insights",
            summary_cards=summary_cards,
            table_section_title="Organization Directory",
            headers=pdf_headers,
            rows=rows,
            status_col_index=6,
            total_count=total_orgs,
            generated_by_username=request.user.username,
            filename=f"organizations_{date_str}.pdf",
        )

        log_action(request.user, "Export", "Organizations")
        return response

    return Response(
        {
            "success": False,
            "message": (
                "Invalid export format. Supported formats " "are csv, excel and pdf."
            ),
        },
        status=status.HTTP_400_BAD_REQUEST,
    )